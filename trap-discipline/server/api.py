"""HTTP API and static file serving (localhost only, token protected)."""
from __future__ import annotations

import asyncio
import csv
import io
import json
import re
import logging
import os
import secrets
import time
from datetime import datetime, timezone

from aiohttp import web

from . import ai as AI
from . import symbols as SYM
from . import learning as L
from . import ai_backups as BK
from . import chart as CH
from . import sms as SMS
from . import config, credentials
from . import game as G
from . import rules as R
from . import sizing as S
from . import stats as ST
from . import bybit as BY
from .bybit import BybitError
from .db import now_ms
from .trades import utc_parts

log = logging.getLogger("trap.api")
TOKEN = secrets.token_urlsafe(24)
ALLOWED_HOSTS = {f"127.0.0.1:{config.PORT}", f"localhost:{config.PORT}"}
TAILNET_HOST = {"value": ""}   # set from settings at startup and on save
# Browser push relays: Chrome/Android (Google), Firefox (Mozilla), Safari/iOS
# (Apple), Edge (Microsoft Windows Push Notification Services)
PUSH_RELAYS = ("fcm.googleapis.com", "android.googleapis.com", "push.services.mozilla.com",
               "push.apple.com", "notify.windows.com")
TAILNET_RE = __import__("re").compile(r"^[a-z0-9](?:[a-z0-9-]{0,62}[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]{0,62}[a-z0-9])?)*\.ts\.net$")

EDITABLE = {"setup", "grade", "followed_plan", "pen_atr", "candles_to_reclaim", "oi_confirmed", "note", "emotion", "tags",
            "stop", "entry", "exit", "qty", "initial_qty", "leverage", "liq_price", "direction", "symbol", "date", "time_utc", "taken"}


def app_of(request: web.Request):
    return request.app["holder"]["app"]


def ok(data=None, **kw):
    return web.json_response({"ok": True, "data": data, **kw}, dumps=lambda o: json.dumps(o, default=str))


def fail(msg: str, status: int = 400, **kw):
    return web.json_response({"ok": False, "error": msg, **kw}, status=status)


@web.middleware
async def security(request: web.Request, handler):
    host = request.headers.get("Host", "")
    allowed = ALLOWED_HOSTS | {f"127.0.0.1:{request.app['port']}", f"localhost:{request.app['port']}"}
    # Opt-in: the one private Tailscale Serve address the owner typed into
    # Settings (so a phone on their own tailnet can enroll for push). The app
    # still listens on 127.0.0.1 only; Tailscale forwards to it locally.
    tn = TAILNET_HOST["value"]
    if tn:
        allowed = allowed | {tn, f"{tn}:443"}
    if host not in allowed:
        return web.Response(status=403, text="Forbidden host")
    if request.path.startswith("/api/"):
        tok = request.headers.get("X-Trap-Token") or request.query.get("token")
        if not tok or not secrets.compare_digest(tok, TOKEN):
            return fail("Missing or bad session token. Reload the page.", 401)
    t0 = time.monotonic()
    resp = await handler(request)
    took = time.monotonic() - t0
    if took > 3 and request.path != "/api/events":
        log.warning("slow request: %s %s took %.1fs", request.method, request.path, took)
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("Referrer-Policy", "no-referrer")
    resp.headers.setdefault("X-Frame-Options", "DENY")
    if request.path in ("/", "/index.html"):
        resp.headers["Content-Security-Policy"] = ("default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; "
                                                   "script-src 'self'; connect-src 'self'; font-src 'self'; frame-ancestors 'none'")
        resp.headers["Cache-Control"] = "no-store"
    elif request.path.startswith(("/static/js/", "/static/css/")):
        # Only app.js carries a version tag; the modules it imports do not. Without this the
        # browser could keep yesterday's copy of a page module next to today's app.js after an
        # update. Revalidating is a tiny 304 from this laptop.
        resp.headers["Cache-Control"] = "no-cache"
    return resp


async def hello(request: web.Request):
    """Lets a newly started copy tell whether this is TRAP, and which copy.
    Only answers the folder and process id on the laptop itself."""
    body = {"app": config.APP_NAME, "version": config.APP_VERSION}
    if request.headers.get("Host", "").split(":")[0] in ("127.0.0.1", "localhost"):
        body.update(folder=str(config.ROOT), pid=os.getpid())
    return web.json_response(body, headers={"Cache-Control": "no-store"})


async def index(request: web.Request):
    html = (config.STATIC_DIR / "index.html").read_text(encoding="utf-8")
    html = html.replace("__TRAP_TOKEN__", TOKEN).replace("__TRAP_VERSION__", config.APP_VERSION)
    return web.Response(text=html, content_type="text/html")


# ---------------------------------------------------------------- state / events
async def state(request):
    a = app_of(request)
    acct = a.account
    return ok({
        "mode": "demo" if a.demo else "live",
        "version": config.APP_VERSION,
        "now": now_ms(),
        "session": a.session_status(),
        "settings": {k: v for k, v in a.settings().items() if k != "thresholds"},
        "thresholds": a.th(),
        "connection": connection_info(a),
        "positions": list(acct.positions.values()) if acct else [],
        "wallet": acct.wallet if acct else None,
        "open_trades": a.db.trades("status='open' AND taken=1"),
        # Every tab asks for this every 8 seconds. It used to carry every unreviewed trade in
        # full (about 1 KB each, forever growing); the pages show five and a count.
        "unreviewed": [{k: t.get(k) for k in ("id", "symbol", "direction", "r", "closed_at")}
                       for t in a.db.trades("status='closed' AND taken=1 AND reviewed_at IS NULL", order="closed_at DESC LIMIT 5")],
        "unreviewed_count": a.db.one("SELECT COUNT(*) AS n FROM trades WHERE status='closed' AND taken=1 AND reviewed_at IS NULL")["n"],
        "game": _game_block(a),
        "coach": a.coach.recent(30),
        "mantra": a.coach.mantra(int(now_ms() // 86_400_000)),
        "watch": a.market.overview() if a.market else [],
    })


_GAME_CACHE: dict = {"key": None, "at": 0.0, "val": None}


def _game_block(a) -> dict:
    """Level, streaks, discipline and quests scan the whole XP and trade history. Reuse the
    answer until something changes (a new XP event or trade edit) or 30 seconds pass."""
    row = a.db.one("SELECT (SELECT COALESCE(MAX(id),0) FROM xp_events) AS xp, (SELECT COUNT(*) FROM trades) AS nt, "
                   "(SELECT COALESCE(MAX(updated_at),0) FROM trades) AS tu, (SELECT COALESCE(MAX(ts),0) FROM checkins) AS ci, "
                   "(SELECT COALESCE(MAX(ts),0) FROM study_progress) AS sp")
    key = (id(a.db), tuple(row.values()), datetime.now(timezone.utc).strftime("%Y-%m-%d"))
    c = _GAME_CACHE
    if c["key"] == key and time.monotonic() - c["at"] < 30:
        return c["val"]
    val = {"level": G.level_for(G.total_xp(a.db)), "streaks": G.streaks(a.db), "discipline": G.discipline_score(a.db),
           "quests": G.quests(a.db)}
    c.update(key=key, at=time.monotonic(), val=val)
    return val


def connection_info(a) -> dict:
    acct = a.account
    pub = a.public_ws
    return {
        "demo": a.demo,
        "connected": bool(acct),
        "status": acct.status if acct else ("demo" if a.demo else "not connected"),
        "error": acct.last_error if acct else None,
        "key": acct.key if acct else None,
        "key_masked": credentials.mask(acct.rest.api_key) if acct else "",
        "private_ws": bool(acct and acct.ws and acct.ws.connected),
        "public_ws": bool(pub and pub.connected) if not a.demo else True,
        "public_error": (pub.last_error if pub else None) or (a.rest.last_error if a.rest else None),
        "storage": credentials.storage_kind(),
        "fee": acct.fee if acct else None,
    }


async def events(request):
    a = app_of(request)
    resp = web.StreamResponse(headers={"Content-Type": "text/event-stream", "Cache-Control": "no-cache", "Connection": "keep-alive"})
    await resp.prepare(request)
    q: asyncio.Queue = asyncio.Queue(maxsize=500)
    a.subscribers.add(q)
    try:
        await resp.write(b"event: hello\ndata: {}\n\n")
        while True:
            try:
                msg = await asyncio.wait_for(q.get(), timeout=15)
                await resp.write(f"data: {json.dumps(msg, default=str)}\n\n".encode())
            except asyncio.TimeoutError:
                await resp.write(b": keepalive\n\n")
    except (ConnectionResetError, asyncio.CancelledError):
        pass
    finally:
        a.subscribers.discard(q)
    return resp


# ---------------------------------------------------------------- market
async def overview(request):
    a = app_of(request)
    return ok(a.market.overview())


_BG_REFRESH: dict[str, asyncio.Task] = {}


async def _quiet_refresh(a, sym: str) -> None:
    try:
        await a.market.refresh(sym)
    except Exception as exc:  # noqa: BLE001  (the market loop retries; the page already has data)
        log.info("background refresh %s: %s", sym, exc)


async def market_symbol(request):
    a = app_of(request)
    sym = request.match_info["symbol"].upper()
    sd = a.market.sd(sym)
    if sd.snapshot and request.query.get("refresh"):
        # Answer at once with what the market loop already has (at most ~20s old) and top it
        # up in the background. Waiting for Bybit here held the chart page for seconds.
        t = _BG_REFRESH.get(sym)
        if t is None or t.done():
            _BG_REFRESH[sym] = asyncio.create_task(_quiet_refresh(a, sym))
    BY.INTERACTIVE.set(True)
    if not sd.snapshot:
        try:
            await a.market.refresh(sym, force=True)
        except BybitError as exc:
            return fail(str(exc), 502)
    return ok(sd.snapshot)


async def chart_data(request):
    """Candles + OI + CVD + funding for the Market Monitor at any timeframe."""
    BY.INTERACTIVE.set(True)      # a person is waiting on this: go ahead of background downloads
    a = app_of(request)
    sym = request.match_info["symbol"].upper()
    tf = request.query.get("tf", "15")
    lb = request.query.get("lookback", "3D")
    if sym not in a.settings()["watchlist"]:
        return fail("Pick a symbol from your watchlist.")
    if not CH.allowed(tf, lb):
        return fail("That timeframe and lookback would load too many bars. Pick a shorter lookback.")
    since = request.query.get("since")
    try:
        data = await a.chart.get(sym, tf, lb, int(since) if since and since.isdigit() else None)
    except BybitError as exc:
        return fail(str(exc), 502)
    except Exception as exc:  # noqa: BLE001  (network hiccups: say so instead of a bare 500)
        log.warning("chart %s %s: %s", sym, tf, exc)
        return fail("Could not load chart history from Bybit. Check the VPN and try Refresh.", 502)
    return ok(data)


async def chart_options(request):
    return ok(CH.options())


# ---------------------------------------------------------------- sizing / checklist
async def size(request):
    a = app_of(request)
    b = await request.json()
    sym = (b.get("symbol") or "").upper()
    snap = a.market.sd(sym).snapshot if (a.market and sym) else {}
    th = a.th()
    thin = sym in a.settings()["thin_assets"]
    equity = float(b.get("equity") or (a.account.wallet.get("equity") if a.account and a.account.wallet else 0) or a.equity_at(None))
    fee = a.settings().get("taker_fee_pct") or th["taker_fee_pct"]
    try:
        res = S.plan_position(
            equity=equity, risk_pct=th["risk_pct"], direction=b.get("direction", "Long"),
            entry=float(b["entry"]), stop=float(b["stop"]), target=float(b["target"]) if b.get("target") else None,
            qty_step=snap.get("qty_step") or 0, min_qty=snap.get("min_qty") or 0, tick=snap.get("tick") or 0,
            taker_fee_pct=fee, slip_pct=th["slippage_pct_thin"] if thin else th["slippage_pct_major"],
            leverage=float(b["leverage"]) if b.get("leverage") else None,
            add1_r=th["addon1_trigger_r"], add1_frac=th["addon1_fraction"], add2_r=th["addon2_trigger_r"], add2_frac=th["addon2_fraction"],
            cost_max_pct=th["cost_pct_of_risk_max"], liq_buffer_min=th["liq_buffer_min"])
    except (KeyError, ValueError, TypeError) as exc:
        return fail(f"Bad input: {exc}")
    res["equity"] = equity
    res["equity_source"] = "Bybit wallet" if (a.account and a.account.wallet and not b.get("equity")) else ("entered" if b.get("equity") else "journal (starting equity + P&L)")
    return ok(res)


def decide_verdict(b: dict, session: dict, th: dict) -> tuple[str, list[str], str | None]:
    """Server-side re-check of the checklist verdict. The UI can suggest; the server decides."""
    reasons = []
    if b.get("recovering"):
        reasons.append("Trading to recover a loss. Do not evaluate the setup.")
    if b.get("disqualifier"):
        reasons.append(f"Hard disqualifier: {b['disqualifier']}.")
    sc = b.get("scores") or {}
    total = sum(int(sc.get(k) or 0) for k in ("c1", "c2", "c3", "c4"))
    if int(sc.get("c2") or 0) == 0:
        reasons.append("C2 scored 0: open interest did not rise through the break. Nobody was recruited.")
    if not b.get("c4_closed"):
        reasons.append("No decisive CLOSE back inside yet (rule 5). Never enter before it.")
    grade = "A" if total >= 7 else "B" if total >= 5 else "C" if total >= 3 else "D"
    if session["status"] in ("stand_down", "caution") or b.get("counter_trend") or b.get("macro"):
        drops = {"A": "B", "B": "C", "C": "D", "D": "D"}
        grade = drops[grade]
        reasons_soft = "Grade dropped one letter for session eligibility."
    else:
        reasons_soft = None
    if grade == "D":
        reasons.append("Grade D: below 3 points, no trade.")
    g = b.get("gates") or {}
    if g.get("range_ok") is None:
        reasons.append("Range gate was never evaluated: enter the range boundaries.")
    elif not g.get("range_ok"):
        reasons.append("Range gate failed.")
    if g.get("cost_ok") is None:
        reasons.append("Cost gate was never evaluated: enter entry and stop.")
    elif not g.get("cost_ok"):
        reasons.append("Cost gate failed.")
    if b.get("third") == "middle":
        reasons.append("Middle third of the range (rule 9).")
    if b.get("trend_block"):
        reasons.append("Counter to a clean 4h trend: only trap in the trend direction.")
    for blk in b.get("personal_blocks") or []:
        reasons.append(str(blk))
    verdict = "NO TRADE" if reasons else "GO"
    return verdict, reasons + ([reasons_soft] if reasons_soft else []), grade


async def checklist(request):
    a = app_of(request)
    b = await request.json()
    sess = a.session_status()
    # canonical names, so "long" or "Spring " can't slip past a personal block
    b["direction"] = next((d for d in ("Long", "Short") if str(b.get("direction") or "").strip().lower() == d.lower()), None)
    b["setup"] = next((n for n in config.SETUP_NAMES if str(b.get("setup") or "").strip().lower() == n.lower()), None)
    b["symbol"] = str(b.get("symbol") or "").strip().upper()
    b["personal_blocks"] = []                      # only the server decides these
    if not b["setup"]:
        b["personal_blocks"].append("Not one of the four setups.")
    if not b["direction"]:
        b["personal_blocks"].append("No direction (Long or Short) was set.")
    verdict, reasons, grade = decide_verdict(b, sess, a.th())
    lctx = {"symbol": b["symbol"]}
    try:
        lctx = a.learning_context(b["symbol"], b["setup"], b["direction"], grade)
        blocks = L.restriction_hits(a.db, lctx)
    except Exception:  # noqa: BLE001  (fail closed: if blocks can't be checked, no GO)
        log.exception("personal blocks")
        blocks = ["Personal blocks could not be checked, so this can't be a GO. Try again."]
    if blocks:
        b["personal_blocks"] = b["personal_blocks"] + blocks
        verdict, reasons, grade = decide_verdict(b, sess, a.th())
    b["server_reasons"] = reasons
    b["session"] = sess
    sz = b.get("sizing") or {}
    cur = a.db.execute(
        "INSERT INTO plans(created_at,symbol,direction,setup,level,entry,stop,target,qty,grade,verdict,status,checklist) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (now_ms(), (b.get("symbol") or "").upper(), b.get("direction"), b.get("setup"), b.get("level"), sz.get("entry"),
         sz.get("stop"), sz.get("target"), sz.get("qty"), grade, verdict, "open", json.dumps(b)))
    pid = cur.lastrowid
    try:
        if b["symbol"] in a.settings()["watchlist"]:
            L.record_checklist(a.db, pid, b, verdict, grade)
    except Exception:  # noqa: BLE001  (the report card must never block the checklist)
        log.exception("record checklist call")
    xp = G.award(a.db, "checklist", f"checklist:{pid}", ref=f"plan:{pid}")
    await a._after_xp()
    try:
        lessons = L.lessons_for(L.edge_profile(a.db.trades()), lctx)
    except Exception:  # noqa: BLE001
        log.exception("lessons")
        lessons = []
    return ok({"id": pid, "verdict": verdict, "grade": grade, "reasons": reasons, "xp": xp, "lessons": lessons})


async def plans(request):
    a = app_of(request)
    rows = a.db.query("SELECT * FROM plans ORDER BY created_at DESC LIMIT 100")
    for r in rows:
        r["checklist"] = json.loads(r["checklist"])
    return ok(rows)


async def plan_skip(request):
    """Log a NO TRADE checklist as a skipped setup in one tap."""
    a = app_of(request)
    pid = int(request.match_info["id"])
    p = a.db.one("SELECT * FROM plans WHERE id=?", (pid,))
    if not p:
        return fail("No such plan", 404)
    chk = json.loads(p["checklist"])
    date, tm = utc_parts(p["created_at"])
    reason = "; ".join(chk.get("server_reasons") or [])[:400]
    tid = a.db.insert_trade({"symbol": p["symbol"], "direction": p["direction"], "taken": 0, "status": "skipped",
                             "source": "manual", "date": date, "time_utc": tm, "setup": p["setup"], "grade": p["grade"],
                             "note": reason or "Skipped", "plan_id": pid, "reviewed_at": now_ms(),
                             "pen_atr": chk.get("pen_atr"), "candles_to_reclaim": chk.get("candles_to_reclaim"),
                             "oi_confirmed": chk.get("oi_confirmed")})
    a.db.execute("UPDATE plans SET trade_id=? WHERE id=?", (tid, pid))
    await a.evaluate(tid)
    return ok({"trade_id": tid})


# ---------------------------------------------------------------- journal
async def trades_list(request):
    a = app_of(request)
    return ok(a.db.trades())


async def trade_get(request):
    a = app_of(request)
    t = a.db.trade(int(request.match_info["id"]))
    if not t:
        return fail("Not found", 404)
    t["executions"] = a.db.query("SELECT * FROM executions WHERE trade_id=? ORDER BY exec_time", (t["id"],))
    t["stop_history"] = a.db.query("SELECT ts, stop_loss, source FROM stop_events WHERE symbol=? AND ts>=? AND ts<=? ORDER BY ts",
                                   (t["symbol"], (t.get("opened_at") or 0) - 60_000, t.get("closed_at") or now_ms()))
    if t.get("plan_id"):
        p = a.db.one("SELECT * FROM plans WHERE id=?", (t["plan_id"],))
        if p:
            p["checklist"] = json.loads(p["checklist"])
        t["plan"] = p
    if t.get("opened_at"):
        cs = await a.candles_between(t["symbol"], t["opened_at"] - 12 * 15 * 60_000, (t.get("closed_at") or now_ms()) + 8 * 15 * 60_000)
        # only show the chart when it actually contains the trade's prices
        if cs and t.get("entry") and min(k["l"] for k in cs) * 0.98 <= t["entry"] <= max(k["h"] for k in cs) * 1.02:
            t["candles"] = cs
    return ok(t)


def _clean(b: dict) -> dict:
    out = {}
    for k, v in b.items():
        if k not in EDITABLE:
            continue
        if k in ("pen_atr", "stop", "entry", "exit", "qty", "initial_qty", "leverage", "liq_price"):
            v = None if v in ("", None) else float(v)
        if k == "candles_to_reclaim":
            v = None if v in ("", None) else int(v)
        if k == "taken":
            v = 1 if v in (1, True, "Yes", "yes") else 0
        if k in ("grade",) and v not in (None, "", "A", "B", "C", "D"):
            continue
        if k in ("followed_plan", "oi_confirmed") and v not in (None, "", "Yes", "No"):
            continue
        out[k] = v
    return out


async def trade_create(request):
    a = app_of(request)
    b = await request.json()
    f = _clean(b)
    if not f.get("symbol"):
        return fail("Symbol is required")
    f["symbol"] = f["symbol"].upper()
    taken = f.get("taken", 1)
    now = now_ms()
    if not f.get("date"):
        f["date"], f["time_utc"] = utc_parts(now)
    f.update({"source": "manual", "status": "closed" if taken else "skipped", "reviewed_at": now})
    if taken:
        try:
            dt = datetime.strptime(f"{f['date']} {f.get('time_utc') or '00:00'}", "%Y-%m-%d %H:%M").replace(tzinfo=timezone.utc)
            f["date"], f["time_utc"] = dt.strftime("%Y-%m-%d"), dt.strftime("%H:%M")
            f["opened_at"] = int(dt.timestamp() * 1000)
            f["closed_at"] = min(f["opened_at"] + 60 * 60_000, now_ms())
        except ValueError:
            return fail("Date must be YYYY-MM-DD and time HH:MM (UTC)")
        f.setdefault("initial_qty", f.get("qty"))
    tid = a.db.insert_trade(f)
    if taken and abs(now - (f.get("opened_at") or 0)) < 30 * 60_000:
        a.attach_context(tid)
    t = await a.evaluate(tid)
    return ok(t)


async def trade_update(request):
    a = app_of(request)
    tid = int(request.match_info["id"])
    t = a.db.trade(tid)
    if not t:
        return fail("Not found", 404)
    b = await request.json()
    f = _clean(b)
    if t["source"] == "bybit":
        for k in ("entry", "exit", "qty", "initial_qty", "symbol", "direction", "date", "time_utc", "taken"):
            f.pop(k, None)  # exchange facts are not editable
        if t.get("stop") and "stop" in f and t.get("status") != "open":
            f.pop("stop")
    if b.get("mark_reviewed") and not t.get("reviewed_at") and t["status"] in ("closed", "skipped"):
        f["reviewed_at"] = now_ms()
    a.db.update_trade(tid, f)
    before_xp = G.total_xp(a.db)
    t = await a.evaluate(tid)
    return ok(t, xp_gained=G.total_xp(a.db) - before_xp)


async def trade_delete(request):
    a = app_of(request)
    tid = int(request.match_info["id"])
    t = a.db.trade(tid)
    if not t:
        return fail("Not found", 404)
    if t["source"] == "bybit":
        return fail("Exchange trades cannot be deleted; they are facts. You can annotate them.", 403)
    a.db.delete_trade(tid)
    for k in ("journal", "fresh", "calib", "clean", "viol", "skip"):
        G.revoke(a.db, f"{k}:{tid}")
    return ok()


async def trade_dismiss(request):
    a = app_of(request)
    tid = int(request.match_info["id"])
    b = await request.json()
    t = a.db.trade(tid)
    if not t:
        return fail("Not found", 404)
    reason = (b.get("reason") or "").strip()
    if len(reason) < 8:
        return fail("Explain in a sentence why this finding is wrong.")
    v = t.get("violations") or []
    for x in v:
        if x["code"] == b.get("code"):
            x["dismissed"] = not b.get("undo")
            x["dismiss_reason"] = reason
    # let evaluate() recompute followed_plan from the surviving findings
    a.db.update_trade(tid, {"violations": v, "followed_plan": None})
    t = await a.evaluate(tid)
    return ok(t)


# ---------------------------------------------------------------- stats / game
async def stats(request):
    a = app_of(request)
    trades = a.db.trades()
    eq0 = a.settings()["starting_equity"]
    th = a.th()
    closed = [t for t in trades if t.get("taken") and t.get("r") is not None]
    closed.sort(key=lambda t: t.get("closed_at") or 0)
    roll = []
    for i in range(len(closed)):
        w = closed[max(0, i - 9): i + 1]
        roll.append({"i": i + 1, "adherence": sum(1 for t in w if t.get("followed_plan") == "Yes") / len(w),
                     "expectancy": sum(t["r"] for t in w) / len(w),
                     "process": sum((t.get("process_score") or 0) for t in w) / len(w)})
    return ok({"kpis": ST.kpis(trades, eq0), "equity": ST.equity_curve(trades, eq0), "diagnostics": ST.diagnostics(trades, th),
               "histogram": ST.r_histogram(trades), "calendar": ST.calendar(trades), "matrix": ST.process_outcome_matrix(trades),
               "violations": ST.violation_counts(trades), "rolling": roll, "thresholds": th,
               "rules": {r["n"]: r["title"] for r in a.rulebook()["rules"]}})


async def game(request):
    a = app_of(request)
    return ok(G.summary(a.db))


async def badges_seen(request):
    a = app_of(request)
    a.db.execute("UPDATE badges SET seen=1")
    return ok()


async def checkin(request):
    a = app_of(request)
    day = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    first = a.db.one("SELECT 1 AS x FROM checkins WHERE day=?", (day,)) is None
    a.db.execute("INSERT OR IGNORE INTO checkins(day, ts) VALUES(?,?)", (day, now_ms()))
    b = await request.json() if request.can_read_body else {}
    if b.get("macro") is not None:
        a.set_setting("macro_day", day if b.get("macro") else None)
    xp = G.award(a.db, "checkin", f"checkin:{day}")
    await a._after_xp()
    return ok({"first": first, "xp": xp, "session": a.session_status()})


# ---------------------------------------------------------------- coach
async def coach_list(request):
    a = app_of(request)
    return ok(a.coach.recent(100))


async def coach_ack(request):
    a = app_of(request)
    b = await request.json()
    a.coach.ack(b.get("id"))
    return ok()


# ---------------------------------------------------------------- playbook / review
def review_state(a) -> dict:
    month = datetime.now(timezone.utc).strftime("%Y-%m")
    a.db.execute("UPDATE reviews SET completed_at=?, notes=? WHERE completed_at IS NULL AND month<>?",
                 (now_ms(), json.dumps({"auto": "expired unfinished"}), month))
    open_r = a.db.one("SELECT * FROM reviews WHERE completed_at IS NULL AND month=? ORDER BY id DESC LIMIT 1", (month,))
    done = a.db.one("SELECT * FROM reviews WHERE month=? AND completed_at IS NOT NULL AND json_extract(notes,'$.auto') IS NULL", (month,))
    return {"month": month, "open": open_r, "done_this_month": bool(done), "can_start": not open_r and not done}


async def playbook(request):
    a = app_of(request)
    hist = a.db.query("SELECT id, ts, note FROM rule_versions ORDER BY id DESC")
    return ok({**a.rulebook(), "history": hist, "review": review_state(a)})


async def review_start(request):
    a = app_of(request)
    rs = review_state(a)
    if not rs["can_start"]:
        return fail("A monthly review is already open or done this month. Rule 11: rule changes happen once a month.", 409)
    cur = a.db.execute("INSERT INTO reviews(month, started_at) VALUES(?,?)", (rs["month"], now_ms()))
    # This is the moment approved-and-queued auto-tune changes are allowed to
    # take effect (rule 11). Each still re-validates before it applies.
    applied = None
    try:
        applied = await a.tuner.apply_queued()
    except Exception:  # noqa: BLE001 - a queue hiccup must not block opening the review
        log.exception("apply_queued")
    return ok({"id": cur.lastrowid, "queued_applied": applied})


async def playbook_save(request):
    a = app_of(request)
    rs = review_state(a)
    if not rs["open"]:
        await a.coach.say("warn", "Not now", "Rule 11: rule changes happen at the monthly review, never during a session.", rule=11,
                          key=f"rule11:{now_ms() // 3_600_000}")
        return fail("Rules are locked. Start the monthly review to change them (rule 11).", 423)
    b = await request.json()
    cur_rb = a.rulebook()
    th = dict(cur_rb["thresholds"])
    for k, v in (b.get("thresholds") or {}).items():
        if k in th:
            th[k] = float(v) if not isinstance(th[k], int) or isinstance(v, float) else int(v)
    rules_ = b.get("rules") or cur_rb["rules"]
    setups = b.get("setups") or cur_rb["setups"]
    from .core import RuleLocked
    try:
        rb = a.commit_rule_version(rules_, th, setups, b.get("note") or "Monthly review edit")
    except RuleLocked as exc:
        return fail(str(exc), 423)
    # every open page picks the new rulebook up live: checklist floors, planner
    # add-on triggers, dashboard bucket edges, session gates
    await a.push("rules", {"id": rb["id"], "note": rb["note"]})
    return ok(rb)


async def review_complete(request):
    a = app_of(request)
    rs = review_state(a)
    if not rs["open"]:
        return fail("No open review", 409)
    b = await request.json()
    a.db.execute("UPDATE reviews SET completed_at=?, notes=? WHERE id=?", (now_ms(), json.dumps(b), rs["open"]["id"]))
    G.award(a.db, "monthly_review", f"review:{rs['month']}")
    await a._after_xp()
    return ok()


# ---------------------------------------------------------------- study
async def study_progress(request):
    a = app_of(request)
    rows = a.db.query("SELECT key, ts, data FROM study_progress")
    return ok({r["key"]: {"ts": r["ts"], **json.loads(r["data"])} for r in rows})


async def study_section(request):
    a = app_of(request)
    b = await request.json()
    key = f"{b['guide']}:{b['section']}"
    a.db.execute("INSERT OR REPLACE INTO study_progress(key, ts, data) VALUES(?,?,?)", (key, now_ms(), json.dumps({"done": True})))
    xp = G.award(a.db, "study_section", f"study:{key}", note=f"Studied {b['section']}")
    await a._after_xp()
    return ok({"xp": xp})


async def study_quiz(request):
    a = app_of(request)
    b = await request.json()
    guide, score, total = b["guide"], int(b["score"]), int(b["total"])
    key = f"{guide}:quiz"
    prev = a.db.one("SELECT data FROM study_progress WHERE key=?", (key,))
    best = max(score, json.loads(prev["data"]).get("best", 0) if prev else 0)
    a.db.execute("INSERT OR REPLACE INTO study_progress(key, ts, data) VALUES(?,?,?)", (key, now_ms(), json.dumps({"best": best, "total": total, "last": score})))
    xp = 0
    if score >= 0.8 * total:
        xp += G.award(a.db, "quiz_pass", f"quiz_pass:{guide}", note=f"Passed the {guide} quiz")
    if score == total:
        xp += G.award(a.db, "quiz_perfect", f"quiz_perfect:{guide}", note=f"Perfect {guide} quiz")
    await a._after_xp()
    return ok({"xp": xp, "best": best})


# ---------------------------------------------------------------- settings / connection
async def settings_get(request):
    a = app_of(request)
    return ok({"settings": a.settings(), "connection": connection_info(a)})


async def settings_put(request):
    a = app_of(request)
    b = await request.json()
    restart = False
    notes: list[str] = []
    if "watchlist" in b:
        raws = [str(x) for x in (b["watchlist"] or []) if x and str(x).strip()]
        if not raws:
            return fail("Watchlist cannot be empty")
        try:
            wl, notes, bad = await SYM.resolve_list(a.rest, raws, demo=a.demo)
        except Exception as exc:  # noqa: BLE001  (Bybit unreachable: say so, change nothing)
            return fail(f"Could not check those coins with Bybit right now ({exc}). Check the VPN and save again.", 502)
        if bad:
            return fail("Bybit has no USDT perpetual for " + ", ".join(bad) +
                        ". Type the coin as it appears on Bybit's derivatives page (for example XRP or XRPUSDT). Nothing was saved.")
        if len(wl) > 8:
            return fail("The watchlist holds up to 8 coins. Remove some and save again.")
        restart = wl != a.settings()["watchlist"]
        a.set_setting("watchlist", wl)
    if "thin_assets" in b:
        thin = []
        for x in b["thin_assets"] or []:
            c = SYM.candidates(x)
            if c:
                # match a watchlist coin however it was typed ("PEPE" -> "1000PEPEUSDT")
                hit = next((w for w in a.settings()["watchlist"] if w in c), c[0])
                if hit not in thin:
                    thin.append(hit)
        a.set_setting("thin_assets", thin)
    for k in ("starting_equity", "backfill_days"):
        if k in b and b[k] not in (None, ""):
            try:
                a.set_setting(k, float(b[k]) if k == "starting_equity" else int(b[k]))
            except (TypeError, ValueError):
                return fail(f"{k.replace('_', ' ')} must be a number")
    for k in ("sound", "notifications", "reduced_motion", "auto_tune"):
        if k in b:
            a.set_setting(k, bool(b[k]))
    if "sms_enabled" in b:
        if bool(b["sms_enabled"]) and not SMS.load():
            return fail("Save a Twilio account under Settings > Notifications first, then turn SMS alerts on.")
        a.set_setting("sms_enabled", bool(b["sms_enabled"]))
    if "tailnet_host" in b:
        tn = str(b["tailnet_host"] or "").strip().lower().removeprefix("https://").rstrip("/")
        if tn and not TAILNET_RE.match(tn):
            return fail("That is not a Tailscale address. It should look like your-laptop.tailXXXX.ts.net")
        a.set_setting("tailnet_host", tn)
        TAILNET_HOST["value"] = tn
    if "push_states" in b:
        if not isinstance(b["push_states"], list):
            return fail("push_states must be a list")
        from .core import PUSH_STATE_CHOICES
        a.set_setting("push_states", [x for x in b["push_states"] if x in PUSH_STATE_CHOICES])
    if "coach_name" in b:
        a.set_setting("coach_name", str(b["coach_name"])[:24] or "Coach")
    if "macro_today" in b:
        a.set_setting("macro_day", datetime.now(timezone.utc).strftime("%Y-%m-%d") if b["macro_today"] else None)
    if restart:
        await a.restart_market()
    if b.get("auto_tune") is True and a.tuner.enabled() and not a.tuner.running \
            and not (a.db.get("tune_last_scan") or {}).get("ok_at"):
        a._spawn(a.tuner.scan())       # first scan right away, then three a day
    return ok(a.settings(), notes=notes)


# ---------------------------------------------------------------- learning
async def learning_page(request):
    a = app_of(request)
    return ok(L.page(a.db, a.db.trades(), a.th(), set(a.settings()["thin_assets"])))


async def learning_comovement(request):
    a = app_of(request)
    return ok(a.co_movement() | {"history": a.co_movement_history()})


async def learning_lessons(request):
    a = app_of(request)
    q = request.query
    sym = (q.get("symbol") or "").upper()
    try:
        ctx = a.learning_context(sym, q.get("setup") or None, q.get("direction") or None, q.get("grade") or None)
        return ok({"lessons": L.lessons_for(L.edge_profile(a.db.trades(), L.validated_keys(a.db)), ctx),
                   "blocks": L.restriction_hits(a.db, ctx), "context": ctx,
                   "exposure": a.exposure_for(sym, q.get("direction") or None)})
    except Exception:  # noqa: BLE001
        log.exception("lessons")
        return ok({"lessons": [], "blocks": [], "context": {}})


async def learning_review(request):
    a = app_of(request)
    rev = L.save_review(a.db, L.weekly_review(a.db, a.db.trades()))
    return ok(L.page(a.db, a.db.trades(), a.th(), set(a.settings()["thin_assets"])) | {"review": rev})


async def learning_restriction(request):
    a = app_of(request)
    try:
        rid = int(request.match_info["id"])
        L.set_restriction(a.db, rid, request.match_info["action"])
    except (KeyError, ValueError):
        return fail("Unknown proposal or action.", 404)
    return ok(L.page(a.db, a.db.trades(), a.th(), set(a.settings()["thin_assets"])))


# ---------------------------------------------------------------- AI analyst (Claude)
async def ai_info(request):
    return ok(app_of(request).ai.info())


async def ai_key_save(request):
    a = app_of(request)
    b = await request.json()
    key = (b.get("api_key") or "").strip()
    if not AI.looks_like_key(key):
        return fail("That does not look like a Claude API key. It starts with sk-ant- (copy it from the Claude Console > Settings > API keys).")
    try:
        await AI.validate_key(a.http, key)
    except AI.AIError as exc:
        return fail(str(exc), 400)
    AI.key_save(key)
    return ok(a.ai.info())


async def ai_backup_key_save(request):
    a = app_of(request)
    prov = request.match_info["provider"]
    if prov not in BK.PROVIDERS:
        return fail("Unknown provider", 404)
    b = await request.json()
    key = (b.get("api_key") or "").strip()
    if not BK.looks_like_key(prov, key):
        where = "platform.openai.com > API keys (it starts with sk-)" if prov == "openai" else "aistudio.google.com > Get API key"
        return fail(f"That does not look like a {BK.PROVIDERS[prov]['label']} key. Copy it from {where}.")
    try:
        await BK.validate(a.http, prov, key)
    except BK.BackupError as exc:
        return fail(str(exc), 400)
    BK.key_save(prov, key)
    return ok(a.ai.info())


async def ai_backup_key_clear(request):
    prov = request.match_info["provider"]
    if prov in BK.PROVIDERS:
        BK.key_clear(prov)
    return ok(app_of(request).ai.info())


async def ai_key_clear(request):
    AI.key_clear()
    return ok(app_of(request).ai.info())


async def ai_settings(request):
    a = app_of(request)
    b = await request.json()
    if "model" in b:
        if b["model"] not in AI.MODELS:
            return fail("Unknown model")
        a.shared.set("ai_model", b["model"])
    for prov in BK.PROVIDERS:
        if f"model_{prov}" in b:
            m = str(b[f"model_{prov}"] or "").strip()
            if m and not re.fullmatch(r"[A-Za-z0-9._\-]{2,60}", m):
                return fail("Model names use letters, numbers, dots and dashes only.")
            a.shared.set(f"ai_model_{prov}", m or None)
    if "backup_order" in b:
        order = [p for p in (b["backup_order"] or []) if p in BK.PROVIDERS]
        a.shared.set("ai_backup_order", order + [p for p in BK.PROVIDERS if p not in order])
    if "budget_usd" in b:
        try:
            v = float(b["budget_usd"])
        except (TypeError, ValueError):
            return fail("Budget must be a number")
        if not 0 <= v <= 1000:
            return fail("Budget must be between $0 and $1,000 a month")
        a.shared.set("ai_budget_usd", v)
    return ok(a.ai.info())


async def ai_scan_start(request):
    a = app_of(request)
    b = await request.json()
    sym = (b.get("symbol") or "").upper()
    if sym not in a.settings()["watchlist"]:
        return fail("Pick a symbol from your watchlist.")
    try:
        return ok(a.ai.start(sym, recovering=bool(b.get("recovering"))))
    except AI.AINotConnected as exc:
        return fail(str(exc), 409, not_connected=True)
    except AI.AIError as exc:
        return fail(str(exc), 409)


async def ai_scan_status(request):
    a = app_of(request)
    return ok(a.ai.status(request.match_info["symbol"].upper()))


# ---------------------------------------------------------------- web push
def _push_view(sub: dict) -> dict:
    """What the Settings page may see: never the subscription's keys."""
    from urllib.parse import urlparse
    host = urlparse(sub["endpoint"]).hostname or ""
    service = ("Apple" if "apple" in host else "Google" if "google" in host
               else "Mozilla" if "mozilla" in host else "Microsoft" if "windows" in host else host)
    return {"endpoint_tail": sub["endpoint"][-16:], "service": service,
            "label": sub.get("label"), "created_at": sub["created_at"], "last_sent_at": sub.get("last_sent_at"),
            "last_error": sub.get("last_error")}


async def push_info(request):
    from . import webpush as WP
    a = app_of(request)
    if a.vapid is None:
        return ok({"available": False, "reason": WP.IMPORT_ERROR or "push library not installed"})
    return ok({"available": True, "public_key": WP.public_key_b64(a.vapid),
               "devices": [_push_view(s) for s in a.shared.push_subs()],
               "states": a.settings()["push_states"]})


def _valid_sub(sub) -> str | None:
    from urllib.parse import urlparse
    if not isinstance(sub, dict):
        return "subscription missing"
    ep = sub.get("endpoint") or ""
    keys = sub.get("keys") or {}
    u = urlparse(ep)
    # A real subscription always points at a browser vendor's push relay.
    # Accepting only those hosts means this endpoint can never be used to make
    # the app send requests to arbitrary local or LAN addresses (server-side
    # request forgery).
    host = (u.hostname or "").lower()
    if u.scheme != "https" or not any(host == d or host.endswith("." + d) for d in PUSH_RELAYS):
        return "not a recognized push service endpoint"
    if not (isinstance(keys.get("p256dh"), str) and isinstance(keys.get("auth"), str)) or len(ep) > 2048:
        return "subscription keys missing"
    return None


async def push_subscribe(request):
    a = app_of(request)
    if a.vapid is None:
        return fail("Push notifications are unavailable on this install.", 503)
    b = await request.json()
    sub = b.get("subscription")
    err = _valid_sub(sub)
    if err:
        return fail(err)
    label = (str(b.get("label") or "").strip()[:40]) or None
    a.shared.push_sub_add(sub["endpoint"], sub["keys"]["p256dh"], sub["keys"]["auth"], label)
    return ok({"devices": [_push_view(s) for s in a.shared.push_subs()]})


async def push_unsubscribe(request):
    a = app_of(request)
    b = await request.json()
    ep = b.get("endpoint")
    tail = b.get("endpoint_tail")
    for s in a.shared.push_subs():
        if (ep and s["endpoint"] == ep) or (tail and s["endpoint"].endswith(tail)):
            a.shared.push_sub_remove(s["endpoint"])
    return ok({"devices": [_push_view(s) for s in a.shared.push_subs()]})


async def push_test(request):
    a = app_of(request)
    if a.vapid is None:
        return fail("Push notifications are unavailable on this install.", 503)
    if not a.shared.push_subs():
        return fail("No devices enrolled yet.")
    res = await a.push_notify(("DEMO · " if a.demo else "") + "TRAP test alert",
                              "If you can read this, setup and entry alerts will reach this device.",
                              tag="trap-test", url="/#/settings")
    devices = [_push_view(s) for s in a.shared.push_subs()]
    if not res["sent"]:
        why = "the device was unsubscribed, so it was removed" if res["removed"] and not res["failed"] else \
              "the push service could not be reached. Check this laptop's internet connection"
        return fail(f"Test not delivered: {why}.", 502, devices=devices, result=res)
    return ok({"devices": devices, "result": res})


# ---------------------------------------------------------------- SMS (Twilio)
def _sms_view(creds: dict | None) -> dict:
    if not creds:
        return {"configured": False}
    return {"configured": True, "storage": SMS.storage_kind(),
            "account_sid_masked": SMS.mask(creds["account_sid"]),
            "from_number": creds["from_number"], "to_number": creds["to_number"]}


async def sms_info(request):
    a = app_of(request)
    v = _sms_view(SMS.load())
    v["enabled"] = a.settings()["sms_enabled"]
    return ok(v)


async def sms_config_save(request):
    a = app_of(request)
    b = await request.json()
    account_sid = (b.get("account_sid") or "").strip()
    auth_token = (b.get("auth_token") or "").strip()
    from_number = (b.get("from_number") or "").strip()
    to_number = (b.get("to_number") or "").strip()
    if not SMS.looks_like_sid(account_sid):
        return fail("That does not look like a Twilio Account SID. It starts with AC and is 34 characters (Twilio Console > Account Info).")
    if not auth_token:
        return fail("Enter the Auth Token from Twilio Console > Account Info (click 'view' to reveal it).")
    if not SMS.looks_like_phone(from_number):
        return fail("The Twilio number must be in international format, for example +15551234567 (find it under Phone Numbers > Manage > Active Numbers).")
    if not SMS.looks_like_phone(to_number):
        return fail("Your own phone number must be in international format, for example +15551234567.")
    try:
        await SMS.validate(a.http, account_sid, auth_token)
    except SMS.SmsError as exc:
        return fail(str(exc), 400)
    SMS.save(account_sid, auth_token, from_number, to_number)
    return ok(_sms_view(SMS.load()) | {"enabled": a.settings()["sms_enabled"]})


async def sms_config_clear(request):
    a = app_of(request)
    SMS.clear()
    a.set_setting("sms_enabled", False)
    return ok(_sms_view(None) | {"enabled": False})


async def sms_test(request):
    a = app_of(request)
    creds = SMS.load()
    if not creds:
        return fail("Save a Twilio account first.", 503)
    ok_sent, err = await SMS.send_sms(a.http, creds,
                                      (("DEMO · " if a.demo else "") + "TRAP test alert - if you got this, "
                                       "setup and entry alerts will reach you by text even with your VPN on."))
    if not ok_sent:
        return fail(f"Test not delivered: {err}", 502)
    return ok({"sent": True})


# ---------------------------------------------------------------- auto-tune
async def tune_state(request):
    a = app_of(request)
    st = a.tuner.status()
    from . import tuner as TU
    st["labels"] = TU.LABELS
    st["bounds"] = TU.BOUNDS
    st["thresholds"] = {k: a.th()[k] for k in TU.TUNED}
    st["position_block"] = a.tuner.position_block()
    st["open_position"] = st["position_block"] is not None
    return ok(st)


async def tune_scan(request):
    """Start a scan in the background and return at once; the page polls
    /api/tune until it finishes (a full scan takes a minute or two)."""
    a = app_of(request)
    if a.tuner.running:
        return ok({"started": False, "status": a.tuner.status()})
    a._spawn(a.tuner.scan(manual=True))
    await asyncio.sleep(0)                # let the scan mark itself running
    return ok({"started": True, "status": a.tuner.status()})


async def _tune_act(a, pid: int, action: str) -> dict:
    fn = {"apply": a.tuner.apply, "dismiss": a.tuner.dismiss, "revert": a.tuner.revert}.get(action)
    if not fn:
        return {"ok": False, "status": 400, "error": "Unknown action."}
    return await fn(pid)


async def tune_action(request):
    a = app_of(request)
    try:
        pid = int(request.match_info["id"])
    except ValueError:
        return fail("Bad proposal id.")
    res = await _tune_act(a, pid, request.match_info["action"])
    if not res.get("ok"):
        return fail(res.get("error", "Not done."), res.get("status", 400))
    return ok({"status": a.tuner.status()})


async def tune_act_signed(request):
    """Apply/Dismiss straight from a phone notification's buttons. No session
    token is available there, so the request carries an HMAC signature that
    only exists inside the (end-to-end encrypted) push message for that one
    proposal and action."""
    a = app_of(request)
    q = request.query
    try:
        pid = int(q.get("id", ""))
    except ValueError:
        return web.json_response({"ok": False, "error": "Bad request."}, status=400)
    action, mode, sig = q.get("a", ""), q.get("m", ""), q.get("sig", "")
    if action not in ("apply", "dismiss") or mode not in ("demo", "live"):
        return web.json_response({"ok": False, "error": "Bad request."}, status=400)
    if mode != ("demo" if a.demo else "live"):
        # checked first: the proposal lives in the other mode's database, so it can't be verified here (nothing changes)
        return web.json_response({"ok": False, "error": "The app is in the other mode now. Open TRAP to review it."}, status=409)
    if not a.tuner.verify(mode, pid, action, sig):
        return web.json_response({"ok": False, "error": "Not authorized."}, status=403)
    res = await _tune_act(a, pid, action)
    if not res.get("ok"):
        return web.json_response({"ok": False, "error": res.get("error")}, status=res.get("status", 400))
    if action == "apply":
        msg = res.get("message") if res.get("queued") else "Applied. Your rulebook is updated."
    else:
        msg = "Dismissed."
    return web.json_response({"ok": True, "message": msg})


async def service_worker(request):
    resp = web.FileResponse(config.STATIC_DIR / "sw.js")
    resp.headers["Content-Type"] = "application/javascript"
    resp.headers["Cache-Control"] = "no-cache"  # so an updated worker is picked up promptly
    return resp


async def manifest(request):
    resp = web.FileResponse(config.STATIC_DIR / "manifest.webmanifest")
    resp.headers["Content-Type"] = "application/manifest+json"
    return resp


async def connect(request):
    a = app_of(request)
    if a.demo:
        return fail("Switch to live mode first (Settings > Mode).", 409)
    b = await request.json()
    key, sec = (b.get("api_key") or "").strip(), (b.get("api_secret") or "").strip()
    if not key or not sec:
        return fail("Paste both the API key and the secret.")
    try:
        ev = await a.connect_account(key, sec, save=True)
    except BybitError as exc:
        msg = exc.msg
        if exc.code == 403:
            msg = "Bybit refused the connection (HTTP 403). Check your VPN is on, then try again."
        elif exc.code in (10003, 10004, 10005):
            msg = f"Bybit rejected the key ({exc.code}: {exc.msg}). Re-copy the key and secret exactly."
        return fail(msg, 400)
    if not ev["ok"]:
        return fail("Key refused: " + " ".join(ev["problems"]), 400, key=ev)
    return ok(ev)


async def disconnect(request):
    a = app_of(request)
    await a.disconnect_account()
    return ok()


async def mode(request):
    b = await request.json()
    holder = request.app["holder"]
    demo = bool(b.get("demo"))
    cur = holder["app"]
    if cur.demo == demo:
        return ok({"mode": "demo" if demo else "live"})
    from .core import TrapApp
    await cur.shutdown()
    new = TrapApp(demo=demo)
    await new.startup()
    holder["app"] = new
    (config.DATA_DIR / "mode.json").write_text(json.dumps({"demo": demo}))
    return ok({"mode": "demo" if demo else "live"})


async def reset_demo(request):
    a = app_of(request)
    if not a.demo:
        return fail("Only the demo database can be reset.", 403)
    for tbl in ("trades", "plans", "xp_events", "badges", "coach", "checkins", "study_progress", "executions", "stop_events"):
        a.db.execute(f"DELETE FROM {tbl}")
    a.db.execute("DELETE FROM kv WHERE key='demo_seeded'")
    from .demo import seed_journal
    seed_journal(a.db, a.settings()["starting_equity"])
    return ok()


# ---------------------------------------------------------------- export
EXPORT_COLS = ["Date", "Time UTC", "Symbol", "Setup", "Grade", "Dir", "Taken?", "Entry", "Stop", "Exit", "Qty", "Leverage",
               "Followed plan?", "Pen. ATR", "Candles to reclaim", "OI confirmed?", "Note / why skipped", "Risk $", "Net P&L", "R"]


def _export_rows(a) -> list[list]:
    rows = []
    for t in sorted(a.db.trades(), key=lambda x: (x.get("date") or "", x.get("time_utc") or "")):
        taken = bool(t.get("taken"))
        rows.append([t.get("date"), t.get("time_utc"), t.get("symbol"), t.get("setup"), t.get("grade"), t.get("direction"),
                     "Yes" if taken else "No", t.get("entry") if taken else None, t.get("stop") if taken else None,
                     t.get("exit") if taken else None, t.get("initial_qty") or t.get("qty") if taken else None,
                     t.get("leverage") if taken else None, t.get("followed_plan") if taken else None, t.get("pen_atr"),
                     t.get("candles_to_reclaim"), t.get("oi_confirmed"), t.get("note"),
                     t.get("risk_usd") if taken else None, t.get("net_pnl") if taken else None, t.get("r") if taken else None])
    return rows


async def export_csv(request):
    a = app_of(request)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPORT_COLS)
    w.writerows(_export_rows(a))
    return web.Response(text=buf.getvalue(), content_type="text/csv",
                        headers={"Content-Disposition": f"attachment; filename=trap-journal-{datetime.now():%Y%m%d}.csv"})


async def export_xlsx(request):
    """Fill the original Trap Journal workbook so its Dashboard and Diagnostics keep working."""
    a = app_of(request)
    import openpyxl
    wb = openpyxl.load_workbook(config.TEMPLATE_XLSX)
    ws = wb["Trade Log"]
    for row in ws.iter_rows(min_row=4, max_row=153, max_col=17):
        for c in row:
            c.value = None
    rows = _export_rows(a)[:150]
    for i, r in enumerate(rows):
        rr = 4 + i
        for j, v in enumerate(r[:17]):
            ws.cell(row=rr, column=j + 1, value=v)
        if r[0]:
            try:
                ws.cell(row=rr, column=1, value=datetime.strptime(r[0], "%Y-%m-%d"))
            except ValueError:
                pass
        # values from the exchange (actual fees) replace the flat-fee formulas for app rows
        if r[6] == "Yes" and r[18] is not None:
            ws.cell(row=rr, column=18, value=r[17])
            ws.cell(row=rr, column=19, value=r[18])
            ws.cell(row=rr, column=20, value=r[19])
    wb["Dashboard"]["B4"] = a.settings()["starting_equity"]
    wb.properties.creator = "Weaver, Pat (Assoc-PHIL-CP)"
    wb.properties.lastModifiedBy = "Weaver, Pat (Assoc-PHIL-CP)"
    out = io.BytesIO()
    wb.save(out)
    return web.Response(body=out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f"attachment; filename=Trap-Journal-{datetime.now():%Y%m%d}.xlsx"})


async def export_json(request):
    a = app_of(request)
    data = {"exported_at": now_ms(), "version": config.APP_VERSION, "trades": a.db.trades(),
            "plans": a.db.query("SELECT * FROM plans"), "xp": a.db.query("SELECT * FROM xp_events"),
            "badges": a.db.query("SELECT * FROM badges"), "rulebook": a.db.query("SELECT * FROM rule_versions"),
            "reviews": a.db.query("SELECT * FROM reviews"), "checkins": a.db.query("SELECT * FROM checkins")}
    return web.Response(text=json.dumps(data, default=str, indent=1), content_type="application/json",
                        headers={"Content-Disposition": f"attachment; filename=trap-backup-{datetime.now():%Y%m%d}.json"})


# ---------------------------------------------------------------- wiring
def build_routes(app: web.Application) -> None:
    r = app.router
    r.add_get("/", index)
    r.add_get("/hello", hello)
    r.add_get("/index.html", index)
    r.add_get("/api/state", state)
    r.add_get("/api/events", events)
    r.add_get("/api/overview", overview)
    r.add_get("/api/market/{symbol}", market_symbol)
    r.add_get("/api/chart/options", chart_options)
    r.add_get("/api/chart/{symbol}", chart_data)
    r.add_post("/api/size", size)
    r.add_post("/api/checklist", checklist)
    r.add_get("/api/plans", plans)
    r.add_post("/api/plans/{id}/skip", plan_skip)
    r.add_get("/api/trades", trades_list)
    r.add_post("/api/trades", trade_create)
    r.add_get("/api/trades/{id}", trade_get)
    r.add_patch("/api/trades/{id}", trade_update)
    r.add_delete("/api/trades/{id}", trade_delete)
    r.add_post("/api/trades/{id}/dismiss", trade_dismiss)
    r.add_get("/api/stats", stats)
    r.add_get("/api/game", game)
    r.add_post("/api/badges/seen", badges_seen)
    r.add_post("/api/checkin", checkin)
    r.add_get("/api/coach", coach_list)
    r.add_post("/api/coach/ack", coach_ack)
    r.add_get("/api/playbook", playbook)
    r.add_put("/api/playbook", playbook_save)
    r.add_post("/api/review/start", review_start)
    r.add_post("/api/review/complete", review_complete)
    r.add_get("/api/study", study_progress)
    r.add_post("/api/study/section", study_section)
    r.add_post("/api/study/quiz", study_quiz)
    r.add_get("/api/settings", settings_get)
    r.add_put("/api/settings", settings_put)
    r.add_post("/api/connect", connect)
    r.add_post("/api/disconnect", disconnect)
    r.add_post("/api/mode", mode)
    r.add_post("/api/demo/reset", reset_demo)
    r.add_get("/api/export.csv", export_csv)
    r.add_get("/api/export.xlsx", export_xlsx)
    r.add_get("/api/export.json", export_json)
    r.add_get("/api/push", push_info)
    r.add_post("/api/push/subscribe", push_subscribe)
    r.add_post("/api/push/unsubscribe", push_unsubscribe)
    r.add_post("/api/push/test", push_test)
    r.add_get("/api/sms", sms_info)
    r.add_post("/api/sms/config", sms_config_save)
    r.add_delete("/api/sms/config", sms_config_clear)
    r.add_post("/api/sms/test", sms_test)
    r.add_get("/api/learning", learning_page)
    r.add_get("/api/learning/comovement", learning_comovement)
    r.add_get("/api/learning/lessons", learning_lessons)
    r.add_post("/api/learning/review", learning_review)
    r.add_post("/api/learning/restrictions/{id}/{action}", learning_restriction)
    r.add_get("/api/ai", ai_info)
    r.add_post("/api/ai/key", ai_key_save)
    r.add_delete("/api/ai/key", ai_key_clear)
    r.add_post("/api/ai/backup/{provider}/key", ai_backup_key_save)
    r.add_delete("/api/ai/backup/{provider}/key", ai_backup_key_clear)
    r.add_put("/api/ai/settings", ai_settings)
    r.add_post("/api/ai/scan", ai_scan_start)
    r.add_get("/api/ai/scan/{symbol}", ai_scan_status)
    r.add_get("/api/tune", tune_state)
    r.add_post("/api/tune/scan", tune_scan)
    r.add_post("/api/tune/{id}/{action}", tune_action)
    r.add_post("/tune-act", tune_act_signed)
    r.add_get("/sw.js", service_worker)
    r.add_get("/manifest.webmanifest", manifest)
    r.add_static("/static", config.STATIC_DIR, show_index=False)
